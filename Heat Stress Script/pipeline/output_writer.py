import os
import pandas as pd
from openpyxl.utils import get_column_letter


def save_result(hr_temp_df, summary_df, participant_sheets, input_file):
    # Output File Name
    folder = os.path.dirname(input_file)

    filename = os.path.basename(input_file)

    filename = os.path.splitext(filename)[0]

    output_file = os.path.join(
        folder,
        f"{filename}_Weather.xlsx"
    )

    # Save Excel
    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        # HR_TEMP Sheet
        hr_temp_df.to_excel(
            writer,
            sheet_name=filename,
            index=False
        )

        # Summary Sheet
        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        for pid, info in participant_sheets.items():
 
            duty_df = info["data"]
            twa = info["twa"]
        
            duty_df.to_excel(
                writer,
                sheet_name=str(pid),
                index=False
            )
        
            worksheet = writer.sheets[str(pid)]

            # Duration column format
            duration_col = None
            
            for cell in worksheet[1]:
                if cell.value == "Duration":
                    duration_col = cell.column_letter
                    break
            
            if duration_col:
                for cell in worksheet[duration_col][1:]:   # Skip header
                    cell.number_format = "[h]:mm:ss.000"
        
            last_row = len(duty_df) + 3
        
            worksheet.cell(row=last_row, column=8).value = "Daily TWA WBGT"
            worksheet.cell(row=last_row, column=9).value = twa

        # Auto Fit Columns
        for sheet in writer.sheets.values():

            for column_cells in sheet.columns:

                length = max(
                    len(str(cell.value))
                    if cell.value is not None
                    else 0
                    for cell in column_cells
                )

                sheet.column_dimensions[
                    get_column_letter(column_cells[0].column)
                ].width = length + 2

    print(f"\nSaved : {output_file}")